#!/usr/bin/env python3
"""Génère le dossier de soutenance Bloc 5 en un seul .docx navigable."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn, nsmap
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "Dossier-Soutenance-Bloc5-M2Shop.docx"

CHAPTERS = [
    ("Partie A — Fiches orales (à lire devant le jury)", ROOT / "docs" / "dossier-oral-simplifie.md"),
    ("Partie B — Correspondance grille Bloc 5", ROOT / "docs" / "correspondance-grille-bloc5.md"),
    ("Partie B — Cartographie et contraintes (C5.1.1 / C5.1.2)", ROOT / "docs" / "C5.1.1-C5.1.2-cartographie-et-contraintes.md"),
    ("Partie B — HLD, LLD et schéma d'architecture (C5.2.1)", ROOT / "docs" / "hld-lld-architecture.md"),
    ("Partie B — Dossier d'architecture et d'exploitation", ROOT / "Rendu-Final-DAE-Complet-M2Shop-Observabilite-Securisee.md"),
    ("Partie B — Matrice de flux réglementaire (C5.1.3)", ROOT / "docs" / "matrice-de-flux.md"),
    ("Partie B — Cahier de recettes et compte-rendu (C5.2.2 / C5.3.3)", ROOT / "docs" / "tests-de-validation.md"),
    ("Partie B — Protocole d'intégration et automatisation (C5.3.1 / C5.3.2)", ROOT / "docs" / "C5.3.1-C5.3.2-protocole-integration.md"),
    ("Partie B — Guide d'utilisation SRE (C5.4.1)", ROOT / "docs" / "C5.4.1-guide-utilisation-sre.md"),
    ("Partie B — Démonstration jury détaillée (C5.2.3)", ROOT / "docs" / "demo-jury.md"),
]

NAVY = RGBColor(0x1B, 0x3A, 0x4B)
TEAL = RGBColor(0x1F, 0x6F, 0x6A)
GRAY = RGBColor(0x44, 0x44, 0x44)
HEADER_BG = "1B3A4B"


def set_run_font(run, name="Calibri", size=11, bold=None, italic=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def shade_cell(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "B0B8C1")
        tcBorders.append(el)
    tcPr.append(tcBorders)


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1)
    run._r.append(instr)
    run._r.append(fld2)


def add_toc(doc):
    """Champ SOMMAIRE Word — à actualiser à l'ouverture (F9 / clic droit)."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    hint = OxmlElement("w:t")
    hint.text = "Clic droit → Mettre à jour les champs  |  ou Ctrl+A puis F9"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(hint)
    run._r.append(fld_end)


def apply_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = GRAY
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    for name, size, color, before, after in (
        ("Heading 1", 18, NAVY, 18, 10),
        ("Heading 2", 14, TEAL, 14, 6),
        ("Heading 3", 12, NAVY, 10, 4),
    ):
        st = styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = color
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True


def configure_sections(doc):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.clear()
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = hp.add_run("M2-Shop  ·  Observabilité sécurisée  ·  Bloc 5")
    set_run_font(r, size=9, color=TEAL, italic=True)

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.clear()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = fp.add_run("Dossier de soutenance  —  page ")
    set_run_font(r1, size=9, color=GRAY)
    add_page_number(fp)
    r2 = fp.add_run("  ·  Confidentiel pédagogique")
    set_run_font(r2, size=9, color=GRAY)


def add_inline(paragraph, text, code=False, size=11):
    """Interprète **gras**, *italique* et `code`."""
    pattern = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")
    pos = 0
    for m in pattern.finditer(text):
        if m.start() > pos:
            run = paragraph.add_run(text[pos : m.start()])
            set_run_font(run, size=size, color=GRAY)
        token = m.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, bold=True, color=NAVY)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name="Consolas", size=size - 1, color=TEAL)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=size, italic=True, color=GRAY)
        pos = m.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, size=size, color=GRAY)


def add_code_block(doc, code):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.4)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F4F6F8")
    shd.set(qn("w:val"), "clear")
    pPr.append(shd)
    run = p.add_run(code.rstrip() + "\n")
    set_run_font(run, name="Consolas", size=9, color=NAVY)


def add_table(doc, rows):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, row in enumerate(rows):
        for j in range(cols):
            cell = table.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            val = row[j] if j < len(row) else ""
            add_inline(p, val, size=9)
            set_cell_border(cell)
            if i == 0:
                shade_cell(cell, HEADER_BG)
                for run in p.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    run.font.size = Pt(9)
    doc.add_paragraph()


def is_table_sep(line):
    s = line.strip().replace(" ", "")
    return bool(re.match(r"^\|?[:\-|]+\|[:\-|]*$", s)) and "-" in s


def parse_table_row(line):
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def looks_like_table_row(line):
    return line.strip().startswith("|") and line.strip().count("|") >= 2


def render_markdown(doc, text, skip_first_h1=True):
    lines = text.replace("\r\n", "\n").split("\n")
    i = 0
    skipped_title = False
    in_code = False
    code_buf = []

    while i < len(lines):
        line = lines[i]

        if line.strip().startswith("```"):
            if in_code:
                add_code_block(doc, "\n".join(code_buf))
                code_buf = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue

        if looks_like_table_row(line):
            rows = [parse_table_row(line)]
            i += 1
            while i < len(lines) and (looks_like_table_row(lines[i]) or is_table_sep(lines[i])):
                if not is_table_sep(lines[i]):
                    rows.append(parse_table_row(lines[i]))
                i += 1
            add_table(doc, rows)
            continue

        stripped = line.strip()
        if not stripped or stripped == "---":
            i += 1
            continue

        img = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)$", stripped)
        if img:
            rel = img.group(2).replace("\\", "/")
            img_path = ROOT / rel if not Path(rel).is_absolute() else Path(rel)
            if not img_path.exists():
                alt = ROOT / "docs" / Path(rel).name
                img_path = alt if alt.exists() else img_path
            if img_path.exists():
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run()
                run.add_picture(str(img_path), width=Cm(16.2))
                cap = doc.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = cap.add_run(img.group(1) or "Figure")
                set_run_font(r, size=9, italic=True, color=TEAL)
            i += 1
            continue

        heading = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if heading:
            level = len(heading.group(1))
            title = heading.group(2).strip()
            if skip_first_h1 and not skipped_title and level == 1:
                skipped_title = True
                i += 1
                continue
            if level == 1:
                doc.add_page_break()
            doc.add_heading(title, level=min(level, 3))
            i += 1
            continue

        if stripped.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            add_inline(p, stripped[2:], size=11)
            for run in p.runs:
                run.italic = True
            i += 1
            continue

        bullet = re.match(r"^[-*]\s+(.*)$", stripped)
        if bullet:
            p = doc.add_paragraph(style="List Bullet")
            p.clear()
            add_inline(p, bullet.group(1))
            i += 1
            continue

        numbered = re.match(r"^\d+\.\s+(.*)$", stripped)
        if numbered:
            p = doc.add_paragraph(style="List Number")
            p.clear()
            add_inline(p, numbered.group(1))
            i += 1
            continue

        p = doc.add_paragraph()
        add_inline(p, stripped)
        i += 1


def add_cover(doc):
    for _ in range(3):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("BLOC 5  —  Systèmes et Réseaux")
    set_run_font(r, size=14, bold=True, color=TEAL)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Dossier de soutenance")
    set_run_font(r, size=28, bold=True, color=NAVY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("M2-Shop  ·  Observabilité sécurisée")
    set_run_font(r, size=16, color=TEAL)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Concevoir et mettre en œuvre l'architecture d'un SI")
    set_run_font(r, size=12, italic=True, color=GRAY)

    doc.add_paragraph()

    meta = [
        ("Référentiel", "ANSSI — Guide d'hygiène informatique / Zone Trust Model"),
        ("Environnement", "3 VMs Debian 12 · Vagrant · Ansible (ansible_local)"),
        ("Compétences", "C5.1.1 à C5.4.1"),
        ("Format", "Partie A : fiches orales  ·  Partie B : dossier complet"),
        ("Recette PoC", "Jalons A–D validés en VM le 7 septembre 2026"),
        ("Dépôt", "github.com/dimitriangely/TP-Mrdecker  ·  Wokrshop-final"),
    ]
    table = doc.add_table(rows=len(meta), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (k, v) in enumerate(meta):
        c0, c1 = table.cell(i, 0), table.cell(i, 1)
        c0.text = k
        c1.text = v
        shade_cell(c0, HEADER_BG)
        for run in c0.paragraphs[0].runs:
            set_run_font(run, size=10, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        for run in c1.paragraphs[0].runs:
            set_run_font(run, size=10, color=NAVY)
        set_cell_border(c0)
        set_cell_border(c1)

    doc.add_paragraph()
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = note.add_run(
        "Document unique enrichi. Partie A = fiches à l'oral. "
        "Partie B = DAE, matrice, recettes, IaC, guide. "
        "Sommaire cliquable : clic droit → Mettre à jour les champs (toute la table)."
    )
    set_run_font(r, size=10, italic=True, color=GRAY)

    doc.add_page_break()

    h = doc.add_heading("Sommaire", level=1)
    add_toc(doc)
    hint = doc.add_paragraph()
    r = hint.add_run(
        "Les titres de ce sommaire sont des liens internes (Ctrl+clic). "
        "S'ils affichent encore le texte d'invite, mettez à jour les champs Word."
    )
    set_run_font(r, size=9, italic=True, color=GRAY)
    doc.add_page_break()


def main():
    doc = Document()
    apply_styles(doc)
    configure_sections(doc)
    add_cover(doc)

    for title, path in CHAPTERS:
        if not path.exists():
            raise FileNotFoundError(path)
        doc.add_heading(title, level=1)
        render_markdown(doc, path.read_text(encoding="utf-8"), skip_first_h1=True)
        doc.add_page_break()

    body = doc.element.body
    last = body[-1]
    if last.tag.endswith("p") and "w:br" in last.xml and 'type="page"' in last.xml:
        body.remove(last)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
